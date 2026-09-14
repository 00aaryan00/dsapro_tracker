import docx
from docx.opc.constants import RELATIONSHIP_TYPE as RT
import json
import sys

def get_hyperlink_targets(doc):
    targets = {}
    for rel in doc.part.rels.values():
        if rel.reltype == RT.HYPERLINK:
            targets[rel.rId] = rel._target
    return targets

def extract_cell_data(cell, link_targets):
    data = {"text": "", "links": []}
    for p in cell.paragraphs:
        for r in p.runs:
            text = r.text
            if text:
                data["text"] += text
        # python-docx doesn't provide a built-in way to iterate over hyperlinks easily,
        # but we can check the paragraph's XML elements
        for element in p._element:
            if element.tag.endswith('hyperlink'):
                rId = element.get('{http://schemas.openxmlformats.org/officeDocument/2006/relationships}id')
                url = link_targets.get(rId, '')
                link_text = "".join([n.text for n in element.iter() if n.text])
                if link_text and url:
                    data["links"].append({"text": link_text, "href": url})
    
    data["text"] = data["text"].strip().replace('\n', ' ')
    return data

def parse_docx(filepath):
    doc = docx.Document(filepath)
    link_targets = get_hyperlink_targets(doc)
    
    result = []
    
    for table in doc.tables:
        for row in table.rows:
            row_data = []
            for cell in row.cells:
                cell_data = extract_cell_data(cell, link_targets)
                row_data.append(cell_data)
            result.append(row_data)
            
    with open("course_data.json", "w", encoding="utf-8") as f:
        json.dump(result, f, indent=2, ensure_ascii=False)

if __name__ == "__main__":
    parse_docx(sys.argv[1])
