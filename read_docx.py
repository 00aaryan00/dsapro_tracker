import docx
import sys

def extract_text(filepath):
    doc = docx.Document(filepath)
    with open("doc_output.txt", "w", encoding="utf-8") as f:
        for i, para in enumerate(doc.paragraphs):
            if para.text.strip():
                f.write(f"P[{i}] {para.text.strip()}\n")
        for t_idx, table in enumerate(doc.tables):
            for r_idx, row in enumerate(table.rows):
                row_data = [cell.text.strip().replace('\n', ' ') for cell in row.cells]
                f.write(f"T[{t_idx}] R[{r_idx}] {' | '.join(row_data)}\n")

if __name__ == "__main__":
    extract_text(sys.argv[1])
